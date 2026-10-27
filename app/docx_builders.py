"""기능별 워드 파일 (PDF와 같은 구성)."""
import re
import random

from docx.shared import Cm, Pt

from .comprehension import spec as cspec
from .docx_export import INK_2, INK_3, Doc
from .grammar import spec as gspec


# ---------------------------------------------------------------- 어법 문제

def grammar_docx(gs, academy):
    qs = [q for q in gs.questions if not gspec.validate(q)]
    doc = Doc(academy, gs.title)
    doc.title(gs.title, gs.target if gs.target and gs.options.get("show_grammar", True) else "")
    p = doc.para(size=9, space_after=8)
    doc.run(p, "반:                    이름:                    점수:", size=9, color=INK_2)
    doc.columns(2)
    C, L = gspec.CIRCLED, gspec.LETTERS
    for n, q in enumerate(qs, 1):
        t = q["type"]
        stem = gspec.STEMS["pick_right" if t == "pick" and q["ask"] == "right" else ("pick_wrong" if t == "pick" else t)]
        p = doc.para(keep=True, space_after=2)
        doc.num(p, n)
        doc.run(p, stem, bold=True, size=9.5)
        if t == "paren":
            doc.box(lambda bp, q=q: (doc.run(bp, q["before"]), doc.run(bp, "("), doc.run(bp, q["options"][0], bold=True),
                                     doc.run(bp, " / "), doc.run(bp, q["options"][1], bold=True), doc.run(bp, ")"), doc.run(bp, q["after"])))
        elif t == "blank":
            doc.box(lambda bp, q=q: doc.run(bp, q["sentence"]))
            for k, c in enumerate(q["choices"]):
                doc.para(f"{C[k]} {c}", hanging_cm=0.5, space_after=0, size=9.5)
            doc.para(space_after=4)
        elif t == "pick":
            for k, c in enumerate(q["choices"]):
                doc.para(f"{C[k]} {c}", hanging_cm=0.5, space_after=0, size=9.5)
            doc.para(space_after=4)
        elif t == "arrange":
            doc.para(q["ko"], size=9.5, space_after=1)
            doc.para("(" + " / ".join(q["given"]) + ")", bold=True, size=9.5, space_after=1)
            if q.get("condition"):
                doc.para("조건: " + q["condition"], size=8.5, color=INK_2)
            doc.blank_line(0)
        elif t == "transform":
            doc.box(lambda bp, q=q: doc.run(bp, q["original"]))
            doc.para("→ " + q["instruction"], size=8.5, color=INK_2)
            doc.blank_line(0)
            doc.blank_line(0)
        else:
            pieces = gspec.place_marks(q["passage"], q["marks"]) or [{"text": q["passage"], "mark": None}]
            sym = C if t == "under5" else L

            def fill(bp, pieces=pieces, sym=sym):
                for pc in pieces:
                    if pc["mark"] is None:
                        doc.run(bp, pc["text"], size=9.5)
                    else:
                        doc.run(bp, sym[pc["mark"]], bold=True, size=8.5)
                        doc.run(bp, pc["text"], size=9.5, underline=True)
            doc.box(fill)
            if t == "count":
                doc.para("   ".join(f"{C[k]} {k + 1}개" for k in range(5)), size=9.5, space_after=6)

    doc.new_section(cols=2)
    doc.title("정답 및 해설", gs.title)
    for n, q in enumerate(qs, 1):
        p = doc.para(space_after=1, keep=True)
        doc.num(p, n)
        doc.run(p, gspec.answer_of(q), bold=True, size=9.5)
        doc.run(p, f"  [{gspec.TYPES[q['type']]}]", size=8, color=INK_3)
        if q.get("explain"):
            doc.para(q["explain"], size=8.8, color=INK_2, indent_cm=0.6, space_after=1)
        if q["type"] in gspec.PASSAGE_TYPES:
            sym = gspec.CIRCLED if q["type"] == "under5" else gspec.LETTERS
            for k, m in enumerate(q["marks"]):
                if m["wrong"]:
                    doc.para(f"{sym[k]} {m['text']} → {m['fix']}" + (f" · {m['why']}" if m.get("why") else ""), size=8.5, color=INK_2, indent_cm=0.6, space_after=1)
        doc.para(space_after=3)
    return doc.response(f"{gs.title}_어법문제.docx")


# ---------------------------------------------------------------- 내용이해 워크북

def comprehension_docx(wb, academy, answer: bool):
    opts = {"start_number": 1, "show_grammar": True, "show_translation": True, "split": [6, 4], **(wb.options or {})}
    doc = Doc(academy, f"{wb.title} · {'답지' if answer else '문제지'}")
    n = int(opts["start_number"])
    for idx, p in enumerate(wb.passages):
        if idx:
            doc.page_break()
        doc.section_head(f"{p.get('title') or f'지문 {idx + 1}'}   {wb.title}" + (" · 정답 및 해설" if answer else ""))
        if not answer:
            head = doc.para(keep=True)
            doc.run(head, f"{n}~{n + cspec.PER_PASSAGE - 1}  ", bold=True, color=doc.brand)
            doc.run(head, "다음 글을 읽고 물음에 답하시오.", bold=True)
            ko = (p.get("passage_ko") or "").strip()
            if ko and opts["show_translation"]:
                a, b = opts["split"]
                t = doc.d.add_table(rows=1, cols=2)
                t.style = "Table Grid"
                doc.fix_widths(t, [17.6 * a / (a + b), 17.6 * b / (a + b)])
                doc.run(t.rows[0].cells[0].paragraphs[0], p["passage"], size=9.5)
                doc.run(t.rows[0].cells[1].paragraphs[0], ko, size=8.5, color=INK_2)
                doc.para()
            else:
                doc.box(lambda bp, p=p: doc.run(bp, p["passage"], size=9.5))
        for key, cnt, instr in cspec.SECTIONS:
            items = p.get(key) or []
            hp = doc.para(keep=True, space_after=2)
            hp.paragraph_format.space_before = Pt(6)
            doc.run(hp, f"{n}~{n + cnt - 1}  ", bold=True, color=doc.brand)
            if answer:
                label = cspec.LABELS[key] + (f" · {p['target_grammar']}" if key == "compose" and p.get("target_grammar") else "")
                doc.run(hp, label, bold=True)
            else:
                doc.run(hp, instr, bold=True, size=9.5)
                if key == "compose" and p.get("target_grammar") and opts["show_grammar"]:
                    doc.run(hp, f"  [목표어법: {p['target_grammar']}]", size=8.5, color=INK_2)
            for i, x in enumerate(items):
                num = n + i
                q = doc.para(hanging_cm=0.7, space_after=1)
                doc.num(q, num)
                if answer:
                    if key.startswith("ox_"):
                        doc.run(q, x["a"] + "  ", bold=True)
                        doc.run(q, x.get("why", ""), size=9, color=INK_2)
                    elif key == "compose":
                        doc.run(q, x["answer"], bold=True)
                        if x.get("point"):
                            doc.run(q, "  → " + x["point"], size=9, color=INK_2)
                        if x.get("source"):
                            doc.run(q, f"  [{x['source']} 문장]", size=8, color=INK_3)
                    else:
                        doc.run(q, f"{x['word']} : ")
                        doc.run(q, x["meaning"], bold=True)
                else:
                    if key.startswith("ox_"):
                        doc.run(q, x["q"] + "  (O / X)")
                    elif key == "compose":
                        doc.run(q, x["ko"])
                        doc.para("(" + " / ".join(x["given"]) + ")", bold=True, indent_cm=0.7, space_after=1)
                        doc.blank_line(0.7)
                    else:
                        doc.run(q, f"{x['word']} : ____________________")
            n += cnt
    return doc.response(f"{wb.title}_{'답지' if answer else '문제지'}.docx")


# ---------------------------------------------------------------- 단어장 · 시험지

def vocab_list_docx(wl, academy, hide: str = "", show_example: bool = True):
    doc = Doc(academy, f"{wl.title} · 단어장")
    doc.title(wl.title, f"단어 {len(wl.words)}개" + (" · 뜻 쓰기 연습" if hide == "meaning" else " · 단어 쓰기 연습" if hide == "word" else ""))
    header = ["#", "단어", "뜻", "유의어 / 반의어"] + (["예문"] if show_example else [])
    widths = [0.9, 3.6, 4.4, 3.2] + ([5.7] if show_example else [])
    rows = []
    for i, w in enumerate(wl.words, 1):
        sa = " / ".join(x for x in (("= " + w["syn"]) if w.get("syn") else "", ("↔ " + w["ant"]) if w.get("ant") else "") if x)
        word = "" if hide == "word" else w["word"] + (f"  {w['pos']}" if w.get("pos") else "")
        meaning = "" if hide == "meaning" else w.get("meaning", "")
        row = [str(i), word, meaning, sa]
        if show_example:
            row.append(w.get("example", "") + ("" if hide == "meaning" or not w.get("example_ko") else "\n" + w["example_ko"]))
        rows.append(row)
    doc.table(header, rows, widths, size=9)
    return doc.response(f"{wl.title}_단어장.docx")


def vocab_test_docx(wl, academy, mode: str, count: int, shuffle: bool, seed: int):
    rng = random.Random(seed)
    pool = [w for w in wl.words if w.get("meaning")]
    picked = rng.sample(pool, min(count, len(pool))) if shuffle else pool[:count]
    qs = []
    for w in picked:
        m = mode if mode != "mix" else rng.choice(["en_ko", "ko_en"])
        qs.append((w["word"], w["meaning"]) if m == "en_ko" else (w["meaning"], w["word"]))
    doc = Doc(academy, f"{wl.title} · 단어 시험")
    doc.title(f"{wl.title} 단어 시험")
    doc.para("반:                    이름:                    점수:        / " + str(len(qs)), size=9, color=INK_2, space_after=6)
    doc.para({"en_ko": "영어 단어의 뜻을 쓰시오.", "ko_en": "우리말 뜻에 맞는 영어 단어를 쓰시오."}.get(mode, "영어는 뜻을, 우리말은 영어 단어를 쓰시오."), bold=True, space_after=6)
    doc.table(["#", "문제", "답", "#", "문제", "답"], _two_cols([(str(i), s, "") for i, (s, _) in enumerate(qs, 1)]), [0.8, 4.2, 3.9, 0.8, 4.2, 3.9], size=9.5)
    doc.new_section(cols=1)
    doc.title("정답", f"{wl.title} 단어 시험")
    doc.table(["#", "문제", "정답", "#", "문제", "정답"], _two_cols([(str(i), s, a) for i, (s, a) in enumerate(qs, 1)]), [0.8, 4.2, 3.9, 0.8, 4.2, 3.9], size=9)
    return doc.response(f"{wl.title}_단어시험.docx")


def _two_cols(items):
    half = (len(items) + 1) // 2
    left, right = items[:half], items[half:]
    return [list(left[i]) + (list(right[i]) if i < len(right) else ["", "", ""]) for i in range(half)]


# ---------------------------------------------------------------- 변형문제

def variant_docx(vs, academy):
    from .grammar.spec import CIRCLED as C, place_marks
    from .variant import spec as vspec

    qs = [q for q in vs.questions if not vspec.validate(q)]
    doc = Doc(academy, vs.title)
    doc.title(vs.title)
    doc.para("반:                    이름:                    점수:", size=9, color=INK_2, space_after=8)
    doc.columns(2)

    def choices(items):
        for k, c in enumerate(items):
            doc.para(f"{C[k]} {c}", hanging_cm=0.5, space_after=0, size=9.2)
        doc.para(space_after=5)

    for n, q in enumerate(qs, 1):
        t = q["type"]
        text = vs.passages[q["p"]]["text"]
        p = doc.para(keep=True, space_after=2)
        doc.num(p, n)
        doc.run(p, vspec.STEMS["match_" + q["ask"]] if t == "match" else vspec.STEMS[t], bold=True, size=9.3)
        doc.run(p, f"  [지문 {q['p'] + 1}]", size=7.5, color=INK_3)
        v = vspec.view(q, text)
        if v:  # 어법·영영풀이·서술형
            def marked(bp, txt):
                for k, part in enumerate(re.split(r"\[u\]|\[/u\]", txt)):
                    if part:
                        doc.run(bp, part, size=9.2, underline=k % 2 == 1)
            if v["passage"]:
                doc.box(lambda bp, v=v: marked(bp, v["passage"]))
            if v["box"]:
                doc.box(lambda bp, v=v: doc.run(bp, (v["box_title"] + "  " if v["box_title"] else "") + v["box"], size=9.2))
            if v["choices"]:
                choices(v["choices"])
            for _ in range(v["lines"]):
                doc.para("_" * 48, size=9, color=INK_3, space_after=2)
            doc.para(space_after=4)
            continue
        if t in ("topic", "title", "gist", "match"):
            doc.box(lambda bp: doc.run(bp, text, size=9.2))
            choices(q["choices"])
        elif t == "blank":
            doc.box(lambda bp, q=q: (doc.run(bp, q["before"], size=9.2), doc.run(bp, "               ", size=9.2, underline=True), doc.run(bp, q["after"], size=9.2)))
            choices(q["choices"])
        elif t == "order":
            doc.box(lambda bp, q=q: doc.run(bp, q["intro"], size=9.2))
            for k in "ABC":
                pp = doc.para(space_after=2)
                doc.run(pp, f"({k}) ", bold=True, size=9.2)
                doc.run(pp, q["parts"][k], size=9.2)
            choices(vspec.ORDER_CHOICES)
        elif t == "insert":
            doc.box(lambda bp, q=q: doc.run(bp, q["given"], size=9.2, bold=True))

            def fill(bp, q=q):
                for i, s in enumerate(q["sentences"]):
                    doc.run(bp, s + " ", size=9.2)
                    if i in q["gaps"]:
                        doc.run(bp, f"( {C[q['gaps'].index(i)]} ) ", bold=True, size=9.2)
            doc.box(fill)
        elif t == "irrelevant":
            def fill(bp, q=q):
                doc.run(bp, q["lead"] + " ", size=9.2)
                for i, s in enumerate(q["numbered"]):
                    doc.run(bp, C[i] + " ", bold=True, size=9.2)
                    doc.run(bp, s + " ", size=9.2)
                doc.run(bp, q.get("tail", ""), size=9.2)
            doc.box(fill)
        elif t == "summary":
            doc.box(lambda bp: doc.run(bp, text, size=9.2))
            doc.para("↓", size=9, color=INK_3)
            doc.box(lambda bp, q=q: doc.run(bp, q["summary"], size=9.2))
            for k, (a, b) in enumerate(q["choices"]):
                doc.para(f"{C[k]}  {a}  ……  {b}", size=9.2, space_after=0)
            doc.para(space_after=5)
        elif t == "vocab":
            pieces = place_marks(q["passage"], q["marks"]) or [{"text": q["passage"], "mark": None}]

            def fill(bp, pieces=pieces):
                for pc in pieces:
                    if pc["mark"] is None:
                        doc.run(bp, pc["text"], size=9.2)
                    else:
                        doc.run(bp, C[pc["mark"]], bold=True, size=8.5)
                        doc.run(bp, pc["text"], size=9.2, underline=True)
            doc.box(fill)

    doc.new_section(cols=2)
    doc.title("정답 및 해설", vs.title)
    for n, q in enumerate(qs, 1):
        p = doc.para(space_after=1, keep=True)
        doc.num(p, n)
        doc.run(p, vspec.answer_text(q), bold=True)
        if q["type"] == "order":
            doc.run(p, " " + vspec.ORDER_CHOICES[q["answer"] - 1])
        doc.run(p, f"  [{vspec.TYPES[q['type']]} · 지문 {q['p'] + 1}]", size=8, color=INK_3)
        if q["type"] == "blank":
            doc.para("빈칸: " + q["phrase"], size=8.8, color=INK_2, indent_cm=0.6, space_after=1)
        if q["type"] in ("grammar", "sa_grammar", "definition"):
            for m in q["marks"]:
                if m["wrong"]:
                    doc.para(f"{m['text']} → {m['fix']}" + (f" · {m['why']}" if m.get("why") else ""), size=8.8, color=INK_2, indent_cm=0.6, space_after=1)
        if q["type"] == "vocab":
            for m in q["marks"]:
                if m["wrong"]:
                    doc.para(f"{m['text']} → {m['fix']}" + (f" · {m['why']}" if m.get("why") else ""), size=8.8, color=INK_2, indent_cm=0.6, space_after=1)
        if q.get("explain"):
            doc.para(q["explain"], size=8.8, color=INK_2, indent_cm=0.6, space_after=1)
        doc.para(space_after=3)
    return doc.response(f"{vs.title}_변형문제.docx")


# ---------------------------------------------------------------- 문장 시험지

def sentence_docx(s, academy, groups):
    from .sentence import spec as sspec

    doc = Doc(academy, s.title)
    doc.title(s.title)
    doc.para("반:                    이름:                    점수:", size=9, color=INK_2, space_after=6)
    n = 0
    for g in groups:
        if len(groups) > 1:
            doc.section_head(g["title"])
        for sec in g["sections"]:
            t = sec["type"]
            hp = doc.para(keep=True, space_after=3)
            hp.paragraph_format.space_before = Pt(8)
            doc.run(hp, f"[{sspec.TYPES[t]}] ", bold=True, color=doc.brand, size=9.5)
            doc.run(hp, sspec.STEMS[t], bold=True, size=9.5)
            for q in sec["items"]:
                n += 1
                p = doc.para(hanging_cm=0.7, space_after=1)
                doc.num(p, n)
                if t == "translate":
                    doc.run(p, q["en"])
                    doc.blank_line(0.7)
                    doc.blank_line(0.7)
                elif t == "compose":
                    doc.run(p, q["ko"])
                    if q.get("hint"):
                        doc.run(p, f"  ({q['hint']})", size=8.8, color=INK_2)
                    doc.blank_line(0.7)
                    doc.blank_line(0.7)
                elif t == "arrange":
                    doc.run(p, q["ko"])
                    doc.para("( " + " / ".join(q["given"]) + " )", bold=True, indent_cm=0.7, space_after=1)
                    doc.blank_line(0.7)
                else:
                    for pc in q["pieces"]:
                        if pc.get("blank"):
                            doc.run(p, f" {pc['hint']}", color=INK_3, size=8.5)
                            doc.run(p, "            ", underline=True)
                            doc.run(p, " ")
                        else:
                            doc.run(p, pc["text"])
                    doc.para(q["ko"], size=8.8, color=INK_2, indent_cm=0.7, space_after=5)
    doc.new_section(cols=1)
    doc.title("정답", s.title)
    n = 0
    for g in groups:
        if len(groups) > 1:
            doc.section_head(g["title"])
        for sec in g["sections"]:
            doc.para(f"[{sspec.TYPES[sec['type']]}]", bold=True, color=doc.brand, size=9.5, space_after=2)
            for q in sec["items"]:
                n += 1
                p = doc.para(hanging_cm=0.7, space_after=1)
                doc.num(p, n)
                if sec["type"] == "translate":
                    doc.run(p, q["ko"], size=9)
                elif sec["type"] == "blank":
                    doc.run(p, ", ".join(pc["blank"] for pc in q["pieces"] if pc.get("blank")), bold=True, size=9)
                    doc.run(p, "  — " + q["en"], size=8.5, color=INK_2)
                else:
                    doc.run(p, q["en"], size=9)
    return doc.response(f"{s.title}_문장시험지.docx")


# ---------------------------------------------------------------- 지문 변형

def paraphrase_docx(s, academy, views, show_compare: bool):
    from docx.shared import RGBColor

    red, green = RGBColor(0x9A, 0x33, 0x24), RGBColor(0x2F, 0x5E, 0x3E)
    doc = Doc(academy, s.title)
    for i, v in enumerate(views):
        if i:
            doc.page_break()
        doc.section_head(v["title"])
        doc.para(v["p"].get("out", ""), size=10.5)
    if show_compare:
        doc.new_section(cols=1)
        doc.title("교사용 · 원문 대조", s.title)
        for i, v in enumerate(views):
            doc.section_head(f"{v['title']} · 강도 {v['strength']} · 변형도 {v['ratio']}%")
            p = doc.para(size=9.5)
            for o in v["ops"]:
                if o.get("sp"):
                    doc.run(p, " ", size=9.5)
                if o["op"] == "equal":
                    doc.run(p, o["text"], size=9.5)
                elif o["op"] == "delete":
                    doc.run(p, o["text"], size=9.5, color=red).font.strike = True
                else:
                    doc.run(p, o["text"], size=9.5, color=green, bold=True, underline=True)
            for n in v["p"].get("notes") or []:
                doc.para("• " + n, size=9, color=INK_2, indent_cm=0.3, space_after=1)
            if v["changes"]:
                doc.para()
                doc.table(["원문", "변형"], [[c["from"], c["to"]] for c in v["changes"]], [8.9, 8.9], size=8.6)
    return doc.response(f"{s.title}_지문변형.docx")
