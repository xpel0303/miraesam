"""워드(.docx) 내려받기 공통 도구.

PDF와 같은 내용을 선생님이 직접 고쳐 쓸 수 있는 워드 파일로 만듭니다.
AI를 다시 부르지 않고 저장된 문항으로 바로 만들기 때문에 누르면 즉시 받아집니다.
글꼴은 나눔고딕(없는 컴퓨터에서는 워드가 비슷한 글꼴로 바꿔 보여 줌), 머리말에 학원 로고·이름.
"""
import base64
import io
import re
from urllib.parse import quote

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from fastapi.responses import Response

FONT = "나눔고딕"
INK_2 = RGBColor(0x5E, 0x5E, 0x5E)
INK_3 = RGBColor(0x8E, 0x8E, 0x8C)
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _rgb(hex_color: str) -> RGBColor:
    h = (hex_color or "#3a3a3a").lstrip("#")
    try:
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except (ValueError, IndexError):
        return RGBColor(0x3A, 0x3A, 0x3A)


def _set_font(run_or_style, size=None):
    f = run_or_style.font
    f.name = FONT
    rpr = (run_or_style.element if hasattr(run_or_style, "element") else run_or_style._element).get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    if size:
        f.size = Pt(size)


class Doc:
    def __init__(self, academy, title: str):
        self.d = Document()
        self.col_cm = 17.8  # 지금 단의 글자 너비 (2단이면 줄어듦)
        self.brand = _rgb(getattr(academy, "brand_color", None))
        normal = self.d.styles["Normal"]
        _set_font(normal, 10)
        normal.paragraph_format.space_after = Pt(2)
        normal.paragraph_format.line_spacing = 1.25
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(1.6)
        sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
        # 머리말: 로고 + 학원 이름 ··· 제목
        hp = sec.header.paragraphs[0]
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(17.8), WD_TAB_ALIGNMENT.RIGHT)
        logo = getattr(academy, "logo", None)
        if logo and logo.startswith("data:image/") and ";base64," in logo and "svg" not in logo[:30]:
            try:
                hp.add_run().add_picture(io.BytesIO(base64.b64decode(logo.split(",", 1)[1])), height=Cm(0.45))
                hp.add_run("  ")
            except Exception:
                pass
        self.run(hp, getattr(academy, "name", ""), size=8, color=INK_3)
        self.run(hp, "\t" + title, size=8, color=INK_3)

    # ---- 글자
    def run(self, p, text, *, bold=False, size=None, color=None, underline=False, italic=False):
        r = p.add_run(text)
        _set_font(r, size)
        r.bold, r.underline, r.italic = bold, underline, italic
        if color is not None:
            r.font.color.rgb = color
        return r

    def para(self, text="", *, bold=False, size=None, color=None, indent_cm=None, hanging_cm=None, space_after=None, keep=False):
        p = self.d.add_paragraph()
        if text:
            self.run(p, text, bold=bold, size=size, color=color)
        pf = p.paragraph_format
        if indent_cm is not None:
            pf.left_indent = Cm(indent_cm)
        if hanging_cm is not None:
            pf.left_indent, pf.first_line_indent = Cm(hanging_cm), Cm(-hanging_cm)
        if space_after is not None:
            pf.space_after = Pt(space_after)
        if keep:
            pf.keep_with_next = True
        return p

    def title(self, text, sub=""):
        self.para(text, bold=True, size=17, space_after=1)
        if sub:
            self.para(sub, size=9, color=INK_2, space_after=6)

    def section_head(self, text):
        p = self.para(text, bold=True, size=11.5, keep=True, space_after=3)
        p.paragraph_format.space_before = Pt(10)
        pbdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"), b.set(qn("w:sz"), "8"), b.set(qn("w:color"), str(self.brand))
        pbdr.append(b)
        p._p.get_or_add_pPr().append(pbdr)
        return p

    def num(self, p, n):
        return self.run(p, f"{n:>2}. " if isinstance(n, int) else f"{n} ", bold=True, color=self.brand)

    def blank_line(self, indent_cm=0.6):
        p = self.para(indent_cm=indent_cm, space_after=6)
        p.paragraph_format.space_before = Pt(10)
        self._border(p, "bottom", "8E8E8C", 4)
        return p

    def _border(self, p, side, color, size):
        ppr = p._p.get_or_add_pPr()
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = OxmlElement("w:pBdr")
            ppr.append(pbdr)
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single"), b.set(qn("w:sz"), str(size)), b.set(qn("w:color"), color)
        pbdr.append(b)

    # ---- 표
    def table(self, header: list[str], rows: list[list[str]], widths_cm: list[float], size=9, bold_first=False):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            self.run(c.paragraphs[0], h, bold=True, size=size)
            self._shade(c, "F2F2F0")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                self.run(cells[i].paragraphs[0], v, size=size, bold=bold_first and i == 0)
        self.fix_widths(t, widths_cm)
        return t

    def fix_widths(self, t, widths_cm):
        """워드·한글·LibreOffice 모두에서 칸 너비가 그대로 보이게 고정."""
        t.autofit = False
        tblpr = t._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblpr.append(layout)
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tblpr.append(tblw)
        tblw.set(qn("w:type"), "dxa"), tblw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w * 567)))
        for row in t.rows:
            for c, w in zip(row.cells, widths_cm):
                c.width = Cm(w)

    def _shade(self, cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
        tcpr.append(shd)

    def box(self, build):
        """테두리 상자 한 칸 표. build(paragraph) 로 내용을 채웁니다."""
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        self.fix_widths(t, [self.col_cm - 0.2])
        p = t.rows[0].cells[0].paragraphs[0]
        build(p)
        self.para(space_after=2)
        return p

    # ---- 단 · 쪽
    def columns(self, n: int):
        sect_pr = self.d.sections[-1]._sectPr
        cols = sect_pr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sect_pr.append(cols)
        cols.set(qn("w:num"), str(n))
        self.col_cm = 17.8 if n == 1 else (17.8 - 0.8 * (n - 1)) / n
        cols.set(qn("w:space"), "454")
        cols.set(qn("w:sep"), "1" if n > 1 else "0")

    def new_section(self, cols: int = 1, new_page: bool = True):
        self.d.add_section(WD_SECTION.NEW_PAGE if new_page else WD_SECTION.CONTINUOUS)
        self.columns(cols)

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def response(self, filename: str) -> Response:
        buf = io.BytesIO()
        self.d.save(buf)
        safe = re.sub(r'[\\/:*?"<>|]', "_", filename)
        return Response(buf.getvalue(), media_type=DOCX_MIME, headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(safe)}"})
